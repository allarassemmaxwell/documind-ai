
from io import BytesIO

from docx import Document
from pypdf import PdfReader


def extract_text(filename: str, content: bytes) -> str:
    """Extract readable text from a PDF or DOCX document."""

    if filename.lower().endswith(".pdf"):
        reader = PdfReader(BytesIO(content))

        if reader.is_encrypted:
            raise ValueError(
                "Password-protected PDFs are not supported."
            )

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    document = Document(BytesIO(content))

    paragraph_text = [
        paragraph.text
        for paragraph in document.paragraphs
    ]

    table_text = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]

    return "\n".join(paragraph_text + table_text)
