
from io import BytesIO

from docx import Document
from reportlab.pdfgen import canvas

from backend.app.document_processor import extract_text


def create_pdf(text: str) -> bytes:
    """Create a small PDF in memory for testing."""
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(100, 750, text)
    pdf.save()
    return buffer.getvalue()


def test_extracts_text_from_pdf():
    content = create_pdf("Hello from PDF")

    text = extract_text("sample.pdf", content)

    assert "Hello from PDF" in text


def test_extracts_text_from_docx_paragraph():
    document = Document()
    document.add_paragraph("Hello from Word")

    buffer = BytesIO()
    document.save(buffer)

    text = extract_text("sample.docx", buffer.getvalue())

    assert "Hello from Word" in text


def test_extracts_text_from_docx_table():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Price"

    buffer = BytesIO()
    document.save(buffer)

    text = extract_text("sample.docx", buffer.getvalue())

    assert "Product" in text
    assert "Price" in text
