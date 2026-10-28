from io import BytesIO

from pypdf import PdfWriter
from docx import Document
from reportlab.pdfgen import canvas
from fastapi.testclient import TestClient
from backend.app.main import app

client = TestClient(app)

def test_health_check():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "healthy"}


def test_rejects_unsupported_file_type():
    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "notes.txt",
                b"Some text",
                "text/plain",
            )
        },
    )

    assert response.status_code == 400
    assert "Only PDF and DOCX" in response.json()["detail"]


def test_rejects_empty_pdf():
    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "empty.pdf",
                b"",
                "application/pdf",
            )
        },
    )

    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()



def test_upload_valid_pdf():
    pdf_buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(pdf_buffer)

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "test.pdf",
                pdf_buffer.getvalue(),
                "application/pdf",
            )
        },
    )

    # A blank PDF has no extractable text.
    # The API should reject it with HTTP 422.
    assert response.status_code == 422
    assert "No readable text" in response.json()["detail"]



def test_upload_valid_docx():
    buffer = BytesIO()

    document = Document()
    document.add_paragraph(
        "DocuMind AI extracts text from Word documents."
    )
    document.save(buffer)
    buffer.seek(0)

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "test.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200

    data = response.json()
    assert data["filename"] == "test.docx"
    assert data["characters_extracted"] > 0
    assert "DocuMind AI extracts text" in data["text_preview"]



def test_rejects_file_over_10_mb():
    large_content = b"x" * (10 * 1024 * 1024 + 1)

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "large.pdf",
                large_content,
                "application/pdf",
            )
        },
    )

    assert response.status_code == 413
    assert "Maximum size is 10 MB" in response.json()["detail"]


def test_rejects_fake_pdf():
    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "fake.pdf",
                b"This is not a real PDF file.",
                "application/pdf",
            )
        },
    )

    assert response.status_code == 400
    assert "content does not match" in response.json()["detail"]


def test_rejects_fake_docx():
    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "fake.docx",
                b"This is not a real Word document.",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 400
    assert "content does not match" in response.json()["detail"]


def test_extracts_text_from_docx_table():
    buffer = BytesIO()

    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Research topic"
    table.cell(0, 1).text = "Artificial Intelligence"
    document.save(buffer)
    buffer.seek(0)

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "research.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200

    extracted_text = response.json()["text_preview"]
    assert "Research topic" in extracted_text
    assert "Artificial Intelligence" in extracted_text



def test_extracts_text_from_multi_page_pdf():
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)

    pdf.drawString(72, 750, "Page one: Literature Review")
    pdf.showPage()

    pdf.drawString(72, 750, "Page two: Research Methods")
    pdf.save()

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "multi_page.pdf",
                buffer.getvalue(),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 200

    data = response.json()
    assert "Page one: Literature Review" in data["text_preview"]
    assert "Page two: Research Methods" in data["text_preview"]
